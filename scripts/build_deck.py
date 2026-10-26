"""Build the short solution-brief draft from actual evaluation output.
Run: uv run --with python-pptx python scripts/build_deck.py
"""

import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

report = json.loads(Path("data/evaluation.json").read_text())
pres = Presentation()
pres.slide_width, pres.slide_height = Inches(13.333), Inches(7.5)
INK, GREEN, PAPER, MUTED = "172C31", "095B4D", "FFFFFF", "53666A"


def box(slide, text, x, y, w, h, size=20, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name, p.font.size = "Arial", Pt(size)
        p.font.bold, p.font.color.rgb = bold, RGBColor.from_string(color)
        p.space_after = Pt(14)
    return shape


def slide(title, subtitle=None, dark=False):
    s = pres.slides.add_slide(pres.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(INK if dark else PAPER)
    box(s, title, 0.65, 0.6, 12, 1.2, 36, PAPER if dark else INK, True)
    if subtitle:
        box(s, subtitle, 0.65, 1.8, 12, 0.8, 18, "BCD6CB" if dark else MUTED)
    box(
        s,
        f"FELLO TAKE-HOME  /  Siddhant Sharma  /  {len(pres.slides):02d}",
        0.65,
        7,
        12,
        0.25,
        10,
        "BCD6CB" if dark else MUTED,
    )
    return s


s = slide("Evidence before outreach.", "A renewal-risk worklist, not another dashboard.", True)
box(
    s,
    "The problem\nA CSM owns about 120 accounts. Signals exist, but attention and lead time are limited.",
    0.65,
    3,
    5.5,
    2.4,
    24,
    PAPER,
)
box(
    s,
    "The decision\nRank the next conversations, show the evidence, and keep customer outreach under human control.",
    7,
    3,
    5.5,
    2.4,
    24,
    PAPER,
)
s.notes_slide.notes_text_frame.text = "Challenge the requested automatic red-score email. Explain why a capacity-limited worklist addresses the actual business decision. This is synthetic, not proof of retention uplift."

s = slide(
    "One codebase. Three services.",
    "Python / FastAPI + Postgres + one worker. No agent framework or message broker.",
)
steps = [
    ("Source ledger", "Canonical aliases\nTwo event clocks"),
    ("Durable jobs", "Atomic enqueue\nLeases + retries"),
    ("Stored scores", "Tenant model\nVersioned evidence"),
    ("CSM worklist", "Review evidence\nPrepare, never send"),
]
for i, (title, body) in enumerate(steps):
    x = 0.65 + i * 3.15
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3), Inches(2.75), Inches(2.15))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string("EAF0EB")
    shape.line.fill.background()
    box(s, title, x + 0.2, 3.25, 2.35, 0.6, 22, GREEN, True)
    box(s, body, x + 0.2, 4, 2.35, 1, 18)
box(
    s,
    "Incoming events trigger processing. Daily reconciliation is only a safety net.",
    0.65,
    5.8,
    12,
    0.75,
    22,
    GREEN,
)
s.notes_slide.notes_text_frame.text = "Walk one late-arriving event through atomic ingestion/enqueue, claim, immutable snapshot and UI. At-least-once execution with idempotent output, not exactly-once execution."

s = slide(
    "Trust is a data contract.",
    "Deterministic definitions and enforced boundaries come before generated prose.",
)
box(
    s,
    "Tenant + CSM boundaries\nPostgres RLS with a non-owner role; account ownership checked in the application.\n\nHistorical correctness\nEvent time AND receipt time must precede the scoring cutoff.",
    0.65,
    2.9,
    5.7,
    3.5,
    22,
)
box(
    s,
    "Inspectable state\nStored inputs, event IDs, model version, feature version and contribution terms.\n\nBounded AI\nOpenRouter selects allowlisted facts. Code renders factual sentences. No email tool.",
    7,
    2.9,
    5.7,
    3.5,
    22,
)
s.notes_slide.notes_text_frame.text = "Explain that LLM output is untrusted. No arbitrary SQL, source prose, customer emails or cross-tenant data enter the request. Contributions are associations, not causal explanations. Unknown never means healthy."

s = slide(
    "Measure the failure, not just the demo.",
    "2024 train → 2025 calibration → 2026 held-out replay. Synthetic results only.",
)
chart_data = CategoryChartData()
tenants = [k for k, v in report["tenants"].items() if v["model_kind"] == "logistic"]
chart_data.categories = tenants
chart_data.add_series("Model AUC", [report["tenants"][k]["auc"] for k in tenants])
chart = s.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.65), Inches(3), Inches(6.6), Inches(3.3), chart_data
).chart
chart.has_legend = False
chart.value_axis.minimum_scale = 0
chart.value_axis.maximum_scale = 1
chart.plots[0].has_data_labels = True
chart.plots[0].data_labels.number_format = "0.00"
box(
    s,
    "No consistent advantage over rules.\n\nHigh recall at capacity 10 is near saturation—not proof of skill.\n\nSparse coverage suppresses displayed probabilities. Small tenant uses explicit rules fallback.",
    7.7,
    2.9,
    4.8,
    3.9,
    22,
)
s.notes_slide.notes_text_frame.text = "Do not claim the model improves retention. Show false positives and touch precision in the JSON report. Repeated weekly flags consume capacity. The tiny tenant has no positive held-out outcomes, so recall/AUC are undefined. There is no statistically established win."

s = slide("A narrow system I can explain.", "Submission draft — release gates remain explicit.", True)
box(
    s,
    "Implemented and exercised\nReproducible data + fitted models\nTenant-safe event processing\nAuthenticated worklist + evidence\nConstrained drafting + labelled fallback",
    0.65,
    3,
    5.8,
    3.2,
    22,
    PAPER,
)
box(
    s,
    "Before submission\nPrivate repository access\nOracle HTTPS deployment\nLive OpenRouter verification\nOwner review and recorded Loom",
    7,
    3,
    5.6,
    3.2,
    22,
    PAPER,
)
s.notes_slide.notes_text_frame.text = "Discuss what fails first at 10x: in-memory bootstrap, connection churn/reconciliation, unbounded ledgers, current-state reads, enterprise identity and provider concurrency. Do not represent remote deployment, GitHub push or a live API call as complete."

out = Path("docs/fello-solution-brief-draft.pptx")
pres.save(out)
check = Presentation(out)
assert len(check.slides) == 5
for s in check.slides:
    for shape in s.shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= check.slide_width
        assert shape.top + shape.height <= check.slide_height
print(
    json.dumps(
        {
            "path": str(out),
            "slides": len(check.slides),
            "status": "content/shape bounds verified; visual render pending (LibreOffice unavailable)",
        }
    )
)
